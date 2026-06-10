import unittest
from io import BytesIO
import sys
import os

# Add root folder to sys.path so we can import app
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import split_text_into_chunks, check_file_signature, normalize_cell_value, process_docx_text, trim_text_for_llm, get_int_env, preflight_input_file, build_ocr_spellcheck_prompt, call_ollama_ocr_spellcheck, apply_ocr_spellcheck, apply_ocr_corrections, normalize_import_type, OCR_IMPORT_TYPES
from tools.benchmark_thuvienphapluat_legal_ocr import normalize_for_scoring

class TestSecurityLimits(unittest.TestCase):
    
    def test_split_text_into_chunks(self):
        # Test case: normal small text
        text = "Hello World\nLine 2"
        chunks = split_text_into_chunks(text, max_chars=50)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0], "Hello World\nLine 2")
        
        # Test case: splitting long text
        text_long = "\n".join([f"This is line {i}" for i in range(20)])
        chunks_long = split_text_into_chunks(text_long, max_chars=100)
        self.assertTrue(len(chunks_long) > 1)
        # Ensure no chunk exceeds 100 characters
        for chunk in chunks_long:
            self.assertTrue(len(chunk) <= 100)

    def test_trim_text_for_llm(self):
        text = "A" * 4000
        trimmed = trim_text_for_llm(text, max_chars=3000)
        self.assertEqual(len(trimmed), 3000)
        self.assertEqual(trimmed, "A" * 3000)
        self.assertEqual(trim_text_for_llm("short text", max_chars=3000), "short text")

    def test_get_int_env(self):
        with unittest.mock.patch.dict(os.environ, {"TEST_INT_ENV": "6000"}, clear=False):
            self.assertEqual(get_int_env("TEST_INT_ENV", 123), 6000)
        with unittest.mock.patch.dict(os.environ, {"TEST_INT_ENV": "invalid"}, clear=False):
            self.assertEqual(get_int_env("TEST_INT_ENV", 123), 123)

    @unittest.mock.patch('app.pdfium.PdfDocument')
    def test_preflight_input_file_pdf_limit(self, mock_pdf_document):
        mock_pdf = unittest.mock.MagicMock()
        mock_pdf.__len__.return_value = 21
        mock_pdf_document.return_value = mock_pdf

        with self.assertRaises(ValueError) as context:
            preflight_input_file("sample.pdf", "sample.pdf")

        self.assertIn("PDF", str(context.exception))
        mock_pdf.close.assert_called()

    @unittest.mock.patch('app.Image.open')
    def test_preflight_input_file_image_limit(self, mock_image_open):
        mock_img = unittest.mock.MagicMock()
        mock_img.size = (200, 200)
        mock_image_open.return_value.__enter__.return_value = mock_img

        with unittest.mock.patch('app.MAX_IMAGE_PIXELS', 10000), unittest.mock.patch('app.MAX_IMAGE_EDGE', 1000):
            with self.assertRaises(ValueError) as context:
                preflight_input_file("sample.png", "sample.png")

        self.assertIn("Ảnh", str(context.exception))
            
    def test_normalize_cell_value_formula_injection(self):
        # Safe string
        self.assertEqual(normalize_cell_value("Hello"), "Hello")
        # Formula strings
        self.assertEqual(normalize_cell_value("=SUM(A1:A5)"), "'=SUM(A1:A5)")
        self.assertEqual(normalize_cell_value("+1+2"), "'+1+2")
        self.assertEqual(normalize_cell_value("@dangerous"), "'@dangerous")
        # Negative numbers (should not be prepended with quote)
        self.assertEqual(normalize_cell_value("-123.45"), "-123.45")
        self.assertEqual(normalize_cell_value("-abc"), "'-abc")

    def test_check_file_signature(self):
        # Valid PNG signature
        png_stream = BytesIO(b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR')
        self.assertTrue(check_file_signature(png_stream, '.png'))
        
        # Invalid PNG signature (e.g. text file renamed)
        bad_png_stream = BytesIO(b'some standard text content here')
        self.assertFalse(check_file_signature(bad_png_stream, '.png'))
        
        # Valid PDF signature
        pdf_stream = BytesIO(b'%PDF-1.4\n%...')
        self.assertTrue(check_file_signature(pdf_stream, '.pdf'))
        
        # Valid DOCX signature
        docx_stream = BytesIO(b'PK\x03\x04\x14\x00\x08\x00\x08\x00')
        self.assertTrue(check_file_signature(docx_stream, '.docx'))

    def test_process_docx_text(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Đây là đoạn văn bản test Word.")
        table = doc.add_table(rows=1, cols=2)
        row = table.rows[0]
        row.cells[0].text = "Cột 1"
        row.cells[1].text = "Cột 2"
        
        test_path = "test_temp_docx.docx"
        doc.save(test_path)
        
        try:
            extracted_text = process_docx_text(test_path)
            self.assertIn("Đây là đoạn văn bản test Word.", extracted_text)
            self.assertIn("Cột 1 | Cột 2", extracted_text)
        finally:
            if os.path.exists(test_path):
                os.remove(test_path)

    def test_process_docx_text_limit(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("A" * 31000)
        
        test_path = "test_temp_docx_limit.docx"
        doc.save(test_path)
        
        try:
            with self.assertRaises(ValueError) as context:
                process_docx_text(test_path)
            self.assertIn("vượt quá giới hạn ký tự", str(context.exception))
        finally:
            if os.path.exists(test_path):
                os.remove(test_path)

    @unittest.mock.patch('app.call_ollama_structured_table')
    def test_build_llm_structured_sheet(self, mock_call):
        from openpyxl import Workbook
        from app import build_llm_structured_sheet
        
        # Configure mock return values using the new multi-table format
        mock_call.return_value = {
            "tables": [
                {
                    "name": "Hóa đơn mua hàng",
                    "headers": ["Tên sản phẩm", "Số lượng", "Đơn giá"],
                    "rows": [
                        {"Tên sản phẩm": "Bút bi", "Số lượng": "10", "Đơn giá": "5000"},
                        {"Tên sản phẩm": "Sổ tay", "Số lượng": "2", "Đơn giá": "25000"}
                    ]
                },
                {
                    "name": "Thông tin nhân sự",
                    "headers": ["Trường thông tin", "Giá trị"],
                    "rows": [
                        {"Trường thông tin": "Họ và tên", "Giá trị": "Nguyễn Văn A"},
                        {"Trường thông tin": "Chức vụ", "Giá trị": "Nhân viên"}
                    ]
                }
            ]
        }
        
        wb = Workbook()
        wb.remove(wb.active) # Remove default sheet
        
        mock_results = [
            {
                "filename": "document.pdf",
                "status": "success",
                "pages": [
                    {"page_number": 1, "text": "Some text content for page 1"},
                    {"page_number": 2, "text": "Some text content for page 2"}
                ]
            }
        ]
        
        build_llm_structured_sheet(wb, mock_results)
        
        # Verify sheets were created successfully
        self.assertIn("Hóa đơn mua hàng", wb.sheetnames)
        self.assertIn("Thông tin nhân sự", wb.sheetnames)
        
        # Verify content of "Hóa đơn mua hàng"
        sheet1 = wb["Hóa đơn mua hàng"]
        # Row 1 headers: Tên file, Tên sản phẩm, Số lượng, Đơn giá
        self.assertEqual([c.value for c in sheet1[1]], ["Tên file", "Tên sản phẩm", "Số lượng", "Đơn giá"])
        # Row 2 data: document.pdf, Bút bi, 10, 5000
        self.assertEqual([c.value for c in sheet1[2]], ["document.pdf", "Bút bi", "10", "5000"])
        
        # Verify content of "Thông tin nhân sự"
        sheet2 = wb["Thông tin nhân sự"]
        self.assertEqual([c.value for c in sheet2[1]], ["Tên file", "Trường thông tin", "Giá trị"])
        self.assertEqual([c.value for c in sheet2[2]], ["document.pdf", "Họ và tên", "Nguyễn Văn A"])

    @unittest.mock.patch('app.requests.post')
    def test_call_ollama_structured_table_payload_limits(self, mock_post):
        from app import call_ollama_structured_table

        mock_response = unittest.mock.Mock()
        mock_response.json.return_value = {
            "message": {
                "content": '{"title":"Bảng","headers":["Cột 1"],"rows":[{"Cột 1":"Giá trị"}],"notes":""}'
            }
        }
        mock_response.raise_for_status.return_value = None
        mock_post.return_value = mock_response

        long_text = "X" * 7000
        result = call_ollama_structured_table("sample.pdf", long_text)

        self.assertEqual(result["title"], "Bảng")
        _, kwargs = mock_post.call_args
        payload = kwargs["json"]
        self.assertIn("think", payload)
        self.assertFalse(payload["think"])
        self.assertEqual(payload["options"]["num_predict"], 6000)

        user_message = payload["messages"][1]["content"]
        self.assertIn("X" * 6000, user_message)
        self.assertNotIn("X" * 6001, user_message)

    def test_spellcheck_prompt_structure(self):
        prompt = build_ocr_spellcheck_prompt("Xin chao")
        self.assertIn("<task>", prompt)
        self.assertIn("Return only exact typo corrections. Do not rewrite text.", prompt)
        self.assertIn("wrong must be copied exactly from the input text.", prompt)
        self.assertIn("Do not include any pair where wrong and correct are identical.", prompt)
        self.assertIn("<input_text><![CDATA[", prompt)
        self.assertIn("Xin chao", prompt)
        self.assertNotIn("clean_text", prompt)

    @unittest.mock.patch('app.requests.post')
    def test_call_ollama_ocr_spellcheck_payload(self, mock_post):
        mock_response = unittest.mock.Mock()
        mock_response.json.return_value = {
            "response": '{"corrections":[{"wrong":"Xin chao","correct":"Xin chào"}]}'
        }
        mock_response.raise_for_status.return_value = None
        mock_post.return_value = mock_response

        result = call_ollama_ocr_spellcheck("Xin chao")
        self.assertEqual(result["corrections"], [{"wrong": "Xin chao", "correct": "Xin chào"}])

        _, kwargs = mock_post.call_args
        self.assertTrue(kwargs["json"]["stream"] is False)
        self.assertFalse(kwargs["json"]["think"])
        self.assertEqual(kwargs["json"]["options"]["temperature"], 0.1)
        self.assertIn("Return only exact typo corrections. Do not rewrite text.", kwargs["json"]["prompt"])
        self.assertIn("Do not include any pair where wrong and correct are identical.", kwargs["json"]["prompt"])
        self.assertEqual(kwargs["json"]["format"]["required"], ["corrections"])
        self.assertNotIn("clean_text", kwargs["json"]["prompt"])

    @unittest.mock.patch('app.requests.post')
    def test_call_ollama_ocr_spellcheck_retries_once_on_invalid_json(self, mock_post):
        first_response = unittest.mock.Mock()
        first_response.json.return_value = {"response": "not valid json"}
        first_response.raise_for_status.return_value = None

        second_response = unittest.mock.Mock()
        second_response.json.return_value = {
            "response": '{"corrections":[{"wrong":"Xin chao","correct":"Xin chào"}]}'
        }
        second_response.raise_for_status.return_value = None

        mock_post.side_effect = [first_response, second_response]

        result = call_ollama_ocr_spellcheck("Xin chao")

        self.assertEqual(result["corrections"], [{"wrong": "Xin chao", "correct": "Xin chào"}])
        self.assertEqual(mock_post.call_count, 2)

    @unittest.mock.patch('app.call_ollama_ocr_spellcheck')
    def test_apply_ocr_spellcheck_chunks_large_text(self, mock_call):
        def _side_effect(chunk, chunk_index=None, chunk_total=None):
            self.assertLessEqual(len(chunk), 6000)
            return {
                "corrections": [],
            }

        mock_call.side_effect = _side_effect

        text = "A" * 20000
        result = apply_ocr_spellcheck(text, source_name="sample.pdf", page_number=1)

        self.assertEqual(result["status"], "success")
        self.assertEqual(result["text"], text)
        self.assertEqual(mock_call.call_count, 4)
        self.assertEqual([len(call.args[0]) for call in mock_call.call_args_list], [6000, 6000, 6000, 2000])

    @unittest.mock.patch('app.call_ollama_ocr_spellcheck')
    def test_apply_ocr_spellcheck_applies_corrections_only(self, mock_call):
        mock_call.return_value = {
            "corrections": [
                {"wrong": "Xin chao", "correct": "Xin chào"},
                {"wrong": "hop dong", "correct": "hợp đồng"},
            ]
        }

        result = apply_ocr_spellcheck("Xin chao ve hop dong nay.", source_name="sample.pdf", page_number=1)

        self.assertEqual(result["status"], "success")
        self.assertEqual(result["text"], "Xin chào ve hợp đồng nay.")
        self.assertEqual(
            result["corrections"],
            [
                {"wrong": "Xin chao", "correct": "Xin chào", "replaced_count": 1},
                {"wrong": "hop dong", "correct": "hợp đồng", "replaced_count": 1},
            ],
        )

    def test_apply_ocr_corrections_filters_broad_and_suspicious(self):
        text = "Xin chao ve hop dong nay. 2026 va https://example.com"
        updated_text, applied = apply_ocr_corrections(
            text,
            [
                {"wrong": "Xin chao", "correct": "Xin chào"},
                {"wrong": "hop dong", "correct": "hợp đồng"},
                {"wrong": "2026", "correct": "2025"},
                {"wrong": "https://example.com", "correct": "https://example.org"},
                {"wrong": "ĐỘC LẬP TỰ DO - HẠNH PHÚC", "correct": "Độc lập - Tự do - Hạnh phúc"},
            ],
        )

        self.assertEqual(updated_text, "Xin chào ve hợp đồng nay. 2026 va https://example.com")
        self.assertEqual(
            applied,
            [
                {"wrong": "Xin chao", "correct": "Xin chào", "replaced_count": 1},
                {"wrong": "hop dong", "correct": "hợp đồng", "replaced_count": 1},
            ],
        )

    def test_relaxed_scoring_normalizes_case_and_punctuation(self):
        left = "ĐỘC LẬP TỰ DO - HẠNH PHÚC"
        right = "Độc lập tự do- hạnh phúc"
        self.assertEqual(normalize_for_scoring(left), normalize_for_scoring(right))

    def test_import_type_aliases_fold_complex_into_clear(self):
        self.assertEqual(normalize_import_type("complex"), "clear")
        self.assertEqual(normalize_import_type("clear"), "clear")
        self.assertIn("complex_llm", OCR_IMPORT_TYPES)
        self.assertTrue(OCR_IMPORT_TYPES["complex_llm"]["llm_postprocess"])

if __name__ == '__main__':
    unittest.main()