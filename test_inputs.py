import os
import sys
import time
import requests

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

def create_test_docx(filename="test_docx.docx"):
    print(f"Creating a test Word file: {filename}...")
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", level=1)
        doc.add_paragraph("Độc lập - Tự do - Hạnh phúc")
        doc.add_heading("TÀI LIỆU KIỂM THỬ TRÍCH XUẤT VĂN BẢN", level=2)
        doc.add_paragraph(
            "Đây là văn bản được sinh tự động để kiểm thử khả năng đọc file Word (.docx) "
            "của hệ thống trích xuất văn bản thông minh. Hệ thống cần đọc nội dung này một cách chính xác."
        )
        
        # Add a table
        table = doc.add_table(rows=2, cols=2)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Tiêu chí'
        hdr_cells[1].text = 'Đánh giá'
        
        row_cells = table.rows[1].cells
        row_cells[0].text = 'Trích xuất Word (.docx)'
        row_cells[1].text = 'Hoạt động tốt và nhanh'
        
        doc.save(filename)
        print("Test Word file created successfully.")
    except Exception as e:
        print(f"Error creating Word file: {e}")
        raise e

def test_ocr_endpoint():
    base_url = "http://127.0.0.1:8092"
    ocr_url = f"{base_url}/ocr"
    health_url = f"{base_url}/health"
    
    # 1. Check health
    print("Checking backend health...")
    try:
        r = requests.get(health_url, timeout=5)
        print(f"Health status: {r.status_code}, Response: {r.json()}")
    except Exception as e:
        print(f"Backend health check failed: {e}. Make sure the Flask server is running.")
        return False
        
    # 2. Check files
    files_to_test = {
        "Image (OCR)": "test_ocr_input.png",
        "Scanned PDF (OCR Fallback)": "25-2025-nd-cp.signed_2pages.pdf",
        "Native PDF (Direct Extraction)": "langchain-paddleocr/tests/data/sample_pdf.pdf",
        "Word (Direct Extraction)": "test_docx.docx"
    }
    
    for name, path in files_to_test.items():
        if not os.path.exists(path):
            print(f"Error: Required test file '{path}' for {name} does not exist!")
            return False
            
    print("\n--- Starting OCR / Text Extraction Tests ---")
    
    all_success = True
    for name, filepath in files_to_test.items():
        print(f"\n[Test Case] Processing {name}: {filepath}")
        start_time = time.time()
        
        try:
            with open(filepath, 'rb') as f:
                files = [('files', (os.path.basename(filepath), f, 'application/octet-stream'))]
                data = {
                    "import_type": "clear",
                    "layout_preserve": "true"
                }
                
                response = requests.post(ocr_url, files=files, data=data, timeout=60)
                
            elapsed = time.time() - start_time
            print(f"HTTP Status: {response.status_code}")
            
            if response.status_code == 200:
                result = response.json()
                print(f"API status field: {result.get('status')}")
                
                ocr_results = result.get("results", [])
                if ocr_results:
                    file_status = ocr_results[0].get("status")
                    print(f"File status: {file_status}")
                    
                    if file_status == "success":
                        pages = ocr_results[0].get("pages", [])
                        print(f"Number of pages extracted: {len(pages)}")
                        text_preview = ""
                        for p in pages:
                            p_text = p.get("text", "")
                            text_preview += f"\n--- Trang {p.get('page_number')} (Độ dài: {len(p_text)}) ---\n"
                            text_preview += p_text[:200] + "..." if len(p_text) > 200 else p_text
                        
                        print("=== Extracted Text Preview ===")
                        print(text_preview)
                        print("==============================")
                        print(f"Extraction {name} SUCCESS. Time taken: {elapsed:.2f} seconds.")
                    else:
                        print(f"Extraction {name} FAILED. Message: {ocr_results[0].get('message')}")
                        all_success = False
                else:
                    print("Extraction FAILED. Empty results list.")
                    all_success = False
            else:
                print(f"Extraction FAILED. Server error: {response.text}")
                all_success = False
                
        except Exception as e:
            print(f"Exception while testing {name}: {e}")
            all_success = False
            
    return all_success

if __name__ == "__main__":
    docx_file = "test_docx.docx"
    create_test_docx(docx_file)
    try:
        success = test_ocr_endpoint()
        if success:
            print("\nAll input formats (Image, Scanned PDF, Native PDF, Word) processed successfully!")
        else:
            print("\nSome tests failed. Please inspect the outputs above.")
    finally:
        if os.path.exists(docx_file):
            os.remove(docx_file)
            print(f"Cleaned up temporary file: {docx_file}")
